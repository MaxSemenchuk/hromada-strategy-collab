#!/usr/bin/env python3
"""MinRegion register of international territorial cooperation agreements.

Source: data.gov.ua dataset «Перелік зареєстрованих угод про міжнародне
територіальне співробітництво» (CMU 357-2025, Law 3668-IX). This is the
international analogue of the domestic PIN registry: named parties, legal
kind, free-text sphere — not sister-city branding and not Interreg grants.

Does NOT fold into v7 combined `score`. Does NOT set known=true.

Usage:
  yarn intl-agreements              # fetch CKAN XLSX + build release
  yarn intl-agreements --offline    # rebuild from data/cache/intl-agreements/
  yarn fetch-intl-agreements        # refresh cache only
"""

from __future__ import annotations

import argparse
import json
import re
import ssl
import sys
import urllib.request
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts" / "analysis"))

from build_interreg_layer import (  # noqa: E402
    load_hromada_index,
    match_partner,
    short_name,
)
from build_partnership_map import COUNTRY_UA_TO_ISO  # noqa: E402
from intl_theme import classify_intl_themes  # noqa: E402

CACHE = ROOT / "data" / "cache" / "intl-agreements"
CACHE_XLSX = CACHE / "register.xlsx"
OUT = ROOT / "data" / "releases" / "intl-agreements.json"
MANIFEST = ROOT / "data" / "releases" / "intl-agreements.manifest.json"
PREVIEW = ROOT / "docs" / "assets" / "intl-agreements-preview.json"

CKAN_PACKAGE = "https://data.gov.ua/api/3/action/package_show?id=6c460279-bac0-45ba-ab71-b255289a49f2"
CKAN_FALLBACK_XLSX = (
    "https://data.gov.ua/dataset/ce370d19-f742-44ce-a075-b8076a27f16b/resource/"
    "604c68ba-3d5b-432b-afd5-649aac1d6033/download/"
    "perelik-zareiestrovanikh-ugod-pro-mizhnarodne-teritorialn-spivrobitnitstvo-15-05-2026.xlsx"
)

UA_HDR = (
    "hromada-strategy-collab/0.1 "
    "(+https://github.com/MaxSemenchuk/hromada-strategy-collab; research cache)"
)

EU27 = {
    "AT", "BE", "BG", "HR", "CY", "CZ", "DK", "EE", "FI", "FR", "DE", "GR",
    "HU", "IE", "IT", "LV", "LT", "LU", "MT", "NL", "PL", "PT", "RO", "SK",
    "SI", "ES", "SE",
}

KIND_LABELS = {
    "cbc": "транскордонне",
    "interterritorial": "міжтериторіальне",
    "cbc_interterritorial": "транскордонне та міжтериторіальне",
    "transnational": "транснаціональне",
    "unspecified": "не зазначено",
}
KIND_LABELS_EN = {
    "cbc": "cross-border",
    "interterritorial": "interterritorial",
    "cbc_interterritorial": "cross-border and interterritorial",
    "transnational": "transnational",
    "unspecified": "unspecified",
}

# 3668 uses long official state names («Французька Республіка») that the
# partnership-map dict (short nominative) misses.
_COUNTRY_STEMS: list[tuple[str, str]] = [
    ("польщ", "PL"),
    ("німеччин", "DE"),
    ("румун", "RO"),
    ("молдов", "MD"),
    ("французьк", "FR"),
    ("франція", "FR"),
    ("італійськ", "IT"),
    ("італія", "IT"),
    ("чеськ", "CZ"),
    ("литовськ", "LT"),
    ("литва", "LT"),
    ("естонськ", "EE"),
    ("естонія", "EE"),
    ("словацьк", "SK"),
    ("латвійськ", "LV"),
    ("латвія", "LV"),
    ("австрі", "AT"),
    ("норвег", "NO"),
    ("угорщ", "HU"),
    ("грузі", "GE"),
    ("турецьк", "TR"),
    ("туреччин", "TR"),
    ("британ", "GB"),
    ("королівство", "GB"),  # last-resort for UK rows; checked after Norway
    ("швед", "SE"),
    ("данськ", "DK"),
    ("данія", "DK"),
    ("фінлянд", "FI"),
    ("нідерланд", "NL"),
    ("голланд", "NL"),
    ("бельг", "BE"),
    ("грец", "GR"),
    ("португал", "PT"),
    ("іспан", "ES"),
    ("словен", "SI"),
    ("хорват", "HR"),
    ("болгар", "BG"),
    ("ірланд", "IE"),
    ("швейцар", "CH"),
    ("люксе", "LU"),
    ("мальт", "MT"),
    ("кіпр", "CY"),
]

OBLAST_ONLY = re.compile(
    r"обласн\w+\s+(?:державн\w+\s+|військов\w+\s+)?адміністрац|"
    r"обласн\w+\s+військов|"
    r"обласн\w+\s+рад",
    re.I,
)
HROMADA_HINT = re.compile(
    r"міськ\w+\s+рад|селищн|сільськ\w+\s+рад|територіальн\w+\s+громад",
    re.I,
)


def _ssl_context() -> ssl.SSLContext:
    try:
        import certifi  # type: ignore

        return ssl.create_default_context(cafile=certifi.where())
    except Exception:
        ctx = ssl.create_default_context()
        try:
            ctx.load_default_certs()
        except Exception:
            pass
        return ctx


def http_get_bytes(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": UA_HDR})
    try:
        with urllib.request.urlopen(req, context=_ssl_context(), timeout=90) as resp:
            return resp.read()
    except Exception:
        import subprocess

        return subprocess.check_output(
            ["curl", "-fsSL", "-A", "Mozilla/5.0", url], timeout=90
        )


def http_get_json(url: str) -> dict:
    raw = http_get_bytes(url)
    return json.loads(raw.decode("utf-8"))


def latest_xlsx_url() -> str:
    try:
        pkg = http_get_json(CKAN_PACKAGE)
        resources = (pkg.get("result") or {}).get("resources") or []
        xlsx = [
            r
            for r in resources
            if str(r.get("format") or "").upper() == "XLSX" and r.get("url")
        ]
        xlsx.sort(key=lambda r: r.get("last_modified") or r.get("created") or "", reverse=True)
        if xlsx:
            return xlsx[0]["url"]
    except Exception as exc:
        print(f"  WARN CKAN package_show failed ({exc}); using fallback URL")
    return CKAN_FALLBACK_XLSX


def fetch_xlsx(force: bool = False) -> Path:
    CACHE.mkdir(parents=True, exist_ok=True)
    if CACHE_XLSX.exists() and not force:
        print(f"Using cached {CACHE_XLSX.relative_to(ROOT)}")
        return CACHE_XLSX
    url = latest_xlsx_url()
    print(f"Fetching {url}")
    CACHE_XLSX.write_bytes(http_get_bytes(url))
    print(f"  wrote {CACHE_XLSX.stat().st_size} bytes")
    return CACHE_XLSX


def kind_id(raw: str | None) -> str:
    t = re.sub(r"\s+", " ", (raw or "").lower())
    if not t.strip():
        return "unspecified"
    cbc = "транскордон" in t
    inter = "міжтериторіальн" in t or "міжрегіональн" in t
    trans = "транснаціональн" in t
    if cbc and inter:
        return "cbc_interterritorial"
    if cbc:
        return "cbc"
    if trans:
        return "transnational"
    if inter:
        return "interterritorial"
    if "міжнародн" in t:
        return "interterritorial"
    return "unspecified"


def country_iso(raw: str | None) -> str | None:
    t = (raw or "").strip().lower()
    if not t:
        return None
    direct = COUNTRY_UA_TO_ISO.get(t)
    if direct:
        return direct
    for name, iso in sorted(COUNTRY_UA_TO_ISO.items(), key=lambda kv: -len(kv[0])):
        if name and name in t:
            return iso
    # Norway before the generic «королівство» stem used for the UK.
    if "норвег" in t:
        return "NO"
    for stem, iso in _COUNTRY_STEMS:
        if stem == "королівство":
            continue
        if stem in t:
            return iso
    if "сполучене королівство" in t or "об'єднане королівство" in t:
        return "GB"
    return None


def ua_level(subject: str | None) -> str:
    s = subject or ""
    if HROMADA_HINT.search(s):
        return "hromada"
    if OBLAST_ONLY.search(s):
        return "oblast"
    return "unknown"


def load_rows(xlsx: Path) -> list[dict]:
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit("openpyxl is required to parse the 3668 register XLSX") from exc

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out: list[dict] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        title = str(row[2] or "").strip()
        if not title or title == "_":
            continue
        out.append(
            {
                "n": row[0],
                "reg": str(row[1] or "").strip() or None,
                "title": title,
                "signed": str(row[3] or "").strip() or None,
                "kind_raw": str(row[4] or "").strip() or None,
                "sphere": str(row[5] or "").strip() or None,
                "ua_region": str(row[6] or "").strip() or None,
                "ua_subject": str(row[7] or "").strip() or None,
                "foreign_state": str(row[8] or "").strip() or None,
                "foreign_atu": str(row[9] or "").strip() or None,
                "foreign_subject": str(row[10] or "").strip() or None,
                "duration": str(row[11] or "").strip() or None,
            }
        )
    wb.close()
    return out


def build_release(rows: list[dict]) -> None:
    _by_name, by_stem, by_town = load_hromada_index()
    by_hromada: dict[str, dict] = {}
    stats: Counter[str] = Counter()
    kind_counts: Counter[str] = Counter()
    country_counts: Counter[str] = Counter()
    unmatched: list[dict] = []

    for raw in rows:
        kid = kind_id(raw["kind_raw"])
        iso = country_iso(raw["foreign_state"])
        themes = classify_intl_themes(raw["title"], raw.get("sphere") or "")
        level = ua_level(raw.get("ua_subject"))
        stats["rows"] += 1
        kind_counts[kid] += 1
        country_counts[iso or "unspecified"] += 1
        if iso in EU27:
            stats["eu"] += 1
        if level == "oblast":
            stats["oblast_level"] += 1
            continue
        row, how = match_partner(
            raw.get("ua_subject") or "",
            None,
            raw.get("ua_region"),
            by_stem,
            by_town,
        )
        stats[how] += 1
        entry = {
            "title": raw["title"],
            "kind_id": kid,
            "kind": KIND_LABELS[kid],
            "kind_en": KIND_LABELS_EN[kid],
            "sphere": raw.get("sphere"),
            "theme_ids": themes,
            "partner_name": raw.get("foreign_atu") or raw.get("foreign_subject"),
            "partner_country": iso,
            "partner_country_ua": raw.get("foreign_state"),
            "signed": raw.get("signed"),
            "duration": raw.get("duration"),
            "source": "law3668",
            "ua_level": level,
            "match": how,
        }
        if not row:
            unmatched.append(
                {
                    "title": raw["title"][:160],
                    "ua_subject": raw.get("ua_subject"),
                    "ua_region": raw.get("ua_region"),
                    "partner_country": iso,
                }
            )
            continue
        code = (row.get("Katottg") or row["Name"]).strip()
        hentry = by_hromada.setdefault(
            code,
            {
                "name": row["Name"],
                "short": short_name(row["Name"]),
                "katottg": row.get("Katottg"),
                "oblast": row.get("Oblast"),
                "agreements": [],
            },
        )
        hentry["agreements"].append(entry)

    hromadas = sorted(
        by_hromada.values(),
        key=lambda h: (-len(h["agreements"]), h["short"]),
    )
    for h in hromadas:
        h["agreement_count"] = len(h["agreements"])
        theme_ids: list[str] = []
        seen: set[str] = set()
        for agr in h["agreements"]:
            for tid in agr.get("theme_ids") or []:
                if tid not in seen:
                    seen.add(tid)
                    theme_ids.append(tid)
        h["theme_ids"] = theme_ids
        h["eu_agreement_count"] = sum(
            1 for a in h["agreements"] if a.get("partner_country") in EU27
        )

    generated = datetime.now(timezone.utc).isoformat()
    payload = {
        "generatedAt": generated,
        "warning": (
            "Law 3668-IX / CMU 357-2025 register of international territorial "
            "cooperation agreements (data.gov.ua). Separate from domestic МСС "
            "(Law 1508-VII), from SKEW/C4C twinning, and from Interreg/keep.eu. "
            "Sphere text is free-form; theme_ids are heuristics. Many rows are "
            "oblast-level or unmatched pre-amalgamation councils. Incomplete vs "
            "the Ministry twincities dashboard. Not folded into matching score; "
            "never known=true."
        ),
        "source": {
            "id": "law3668",
            "name": "Перелік зареєстрованих угод про міжнародне територіальне співробітництво",
            "url": "https://data.gov.ua/dataset/6c460279-bac0-45ba-ab71-b255289a49f2",
            "law": "3668-IX",
            "cmu": "357-2025",
        },
        "coverage": {
            "rows": stats["rows"],
            "eu_rows": stats["eu"],
            "oblast_level_skipped": stats["oblast_level"],
            "hromadas_matched": len(hromadas),
            "agreements_matched": sum(h["agreement_count"] for h in hromadas),
            "unmatched": stats.get("unmatched", 0),
            "kind_breakdown": dict(kind_counts.most_common()),
            "country_breakdown": dict(country_counts.most_common()),
            "resolve_stats": {
                k: v
                for k, v in stats.items()
                if k in ("name_stem", "town", "unmatched", "oblast_level", "rows", "eu")
            },
        },
        "hromadas": hromadas,
        "unmatched": unmatched[:80],
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    MANIFEST.write_text(
        json.dumps(
            {
                "generatedAt": generated,
                "rows": stats["rows"],
                "hromadasMatched": len(hromadas),
                "agreementsMatched": payload["coverage"]["agreements_matched"],
                "euRows": stats["eu"],
                "method": "CKAN XLSX Law 3668-IX; name-stem join to hromadas.json; intl_theme on title+sphere",
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    PREVIEW.write_text(
        json.dumps(
            {
                "generatedAt": generated,
                "caveat": payload["warning"],
                "hromadaCount": len(hromadas),
                "agreementsMatched": payload["coverage"]["agreements_matched"],
                "euRows": stats["eu"],
                "top": [
                    {
                        "short": h["short"],
                        "oblast": h["oblast"],
                        "agreement_count": h["agreement_count"],
                        "themes": h["theme_ids"][:6],
                        "partners": [
                            f"{(a.get('partner_name') or '—')} ({a.get('partner_country') or '?'})"
                            for a in h["agreements"][:5]
                        ],
                    }
                    for h in hromadas[:12]
                ],
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print(
        f"Wrote {OUT.relative_to(ROOT)} — rows={stats['rows']} "
        f"matched_hromadas={len(hromadas)} "
        f"matched_agreements={payload['coverage']['agreements_matched']} "
        f"oblast_skipped={stats['oblast_level']} unmatched={stats.get('unmatched', 0)} "
        f"eu={stats['eu']}"
    )


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--offline", action="store_true")
    ap.add_argument("--fetch-only", action="store_true")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    if args.fetch_only:
        fetch_xlsx(force=True)
        return
    if args.offline:
        if not CACHE_XLSX.exists():
            raise SystemExit(f"No cached {CACHE_XLSX} — run without --offline first")
        xlsx = CACHE_XLSX
    else:
        xlsx = fetch_xlsx(force=args.force)
    rows = load_rows(xlsx)
    print(f"Parsed {len(rows)} filled agreement rows")
    build_release(rows)


if __name__ == "__main__":
    main()
