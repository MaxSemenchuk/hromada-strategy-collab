"""UA–EU twinning vs domestic ММС: who participates, how they overlap, what about.

Reproduces every number in internal/twinning-vs-mms-article-draft.md.

Inputs (run from repo root):
  data/releases/{hromadas,twinning-partners,intl-agreements,hromada-resources,edem-barometer}.json
  data/cache/kse/partnerships-hromadas-network.csv   (KSE; ММС registry, flag active_2402)
  data/cache/mss/mss_registry.xlsx                    (Мінрегіон registry, titles + forms)

Definitions
  twinning  = hromada has >=1 partner with confidence == "registry"
              (SKEW, AFCCRE, decentralization.ua Ministry map)
  ММС       = hromada appears in >=1 registry agreement active on 24.02.2022
  Both are registry snapshots, but the ММС flag is pre-war while twinning
  includes post-2022 solidarity partnerships — a timing asymmetry.

Usage: python3 scripts/analysis/twinning_vs_mms.py   (writes internal/twinning-vs-mms-results.json)
"""

import csv
import json
import math
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from intl_theme import classify_intl_themes  # noqa: E402

REL = ROOT / "data" / "releases"
OUT = ROOT / "internal" / "twinning-vs-mms-results.json"
TIERS = ["<5k", "5-10k", "10-20k", "20-50k", "50k+"]


def load(name):
    return json.load(open(REL / name, encoding="utf-8"))


def pop_tier(p):
    if p is None:
        return None
    return "<5k" if p < 5000 else "5-10k" if p < 10000 else "10-20k" if p < 20000 else "20-50k" if p < 50000 else "50k+"


def wilson(k, n, z=1.96):
    if n == 0:
        return None
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return [round(100 * (c - h), 1), round(100 * (c + h), 1)]


def phi(a, b, base):
    n = len(base)
    A, B, AB = len(a & base), len(b & base), len(a & b & base)
    den = math.sqrt(A * (n - A) * B * (n - B))
    return round((AB * n - A * B) / den, 3) if den else None


def share(sets, groups):
    """groups: {label: set of katottg}. Returns {set_name: {label: [k, n, pct, ci]}}."""
    res = {}
    for name, s in sets.items():
        res[name] = {}
        for label, members in groups.items():
            k = len(members & s)
            n = len(members)
            res[name][label] = [k, n, round(100 * k / n, 1) if n else None, wilson(k, n)]
    return res


def main():
    H = load("hromadas.json")
    U = {r["Katottg"]: r for r in H}
    base = set(U)

    T = load("twinning-partners.json")["hromadas"]
    tw = {h["katottg"] for h in T if any(p["confidence"] == "registry" for p in h["partners"])} & base
    tw_src = defaultdict(set)
    for h in T:
        for p in h["partners"]:
            if p["confidence"] == "registry":
                tw_src[p["source"]].add(h["katottg"])

    rows = list(csv.DictReader(open(ROOT / "data/cache/kse/partnerships-hromadas-network.csv", encoding="utf-8")))
    mms = set()
    for r in rows:
        if r["active_2402"] == "1":
            mms |= {r["hromada_code.x"], r["hromada_code.y"]}
    mms &= base
    sets = {"twinning": tw, "mms": mms}

    out = {"universe": len(base), "twinning": len(tw), "mms": len(mms),
           "twinning_by_source": {k: len(v & base) for k, v in tw_src.items()},
           "mms_agreements_active_2402": len({r["register_number"] for r in rows if r["active_2402"] == "1"})}

    both = tw & mms
    out["overlap"] = {"both": len(both), "twinning_only": len(tw - mms), "mms_only": len(mms - tw),
                      "neither": len(base - tw - mms), "phi": phi(tw, mms, base),
                      "share_of_mms_also_twinned": round(100 * len(both) / len(mms), 1)}

    # overlap controlling for size: phi within each population tier
    by_tier = defaultdict(set)
    for k, r in U.items():
        t = pop_tier(r["Population"])
        if t:
            by_tier[t].add(k)
    out["phi_within_tier"] = {t: phi(tw, mms, by_tier[t]) for t in TIERS}

    out["by_population"] = share(sets, {t: by_tier[t] for t in TIERS})
    out["by_type"] = share(sets, {t: {k for k, r in U.items() if r["Type"] == t} for t in ("міська", "селищна", "сільська")})
    out["by_oblast"] = share(sets, {o: {k for k, r in U.items() if r["Oblast"] == o} for o in sorted({r["Oblast"] for r in H})})

    R = {h["katottg"]: h for h in load("hromada-resources.json")["hromadas"]}
    vals = sorted(h["own_income_per_capita"] for h in R.values() if h.get("own_income_per_capita"))
    q = [vals[len(vals) * i // 4] for i in (1, 2, 3)]
    groups = defaultdict(set)
    for k, h in R.items():
        v = h.get("own_income_per_capita")
        if v and k in base:
            groups["Q1" if v < q[0] else "Q2" if v < q[1] else "Q3" if v < q[2] else "Q4"].add(k)
    out["by_own_income_quartile"] = share(sets, dict(sorted(groups.items())))

    E = {h["katottg"]: h for h in load("edem-barometer.json")["hromadas"]}
    groups = defaultdict(set)
    for k, h in E.items():
        if k in base:
            groups[h["tier"]].add(k)
    out["by_edem_tier"] = share(sets, dict(groups))

    donors = {k for k, r in U.items() if r["DonorsPrograms"]}
    out["by_any_donor"] = share(sets, {"donor": donors, "no_donor": base - donors})
    per_prog = {}
    progs = Counter(p for r in H for p in r["DonorsPrograms"])
    for p in progs:
        per_prog[p] = share(sets, {p: {k for k, r in U.items() if p in r["DonorsPrograms"]}})
    out["by_donor_program"] = {p: {s: v[p] for s, v in d.items()} for p, d in per_prog.items()}
    giz = {k for k, r in U.items() if "GIZ" in r["DonorsPrograms"]}
    out["giz_bridge"] = {"giz_hromadas": len(giz), "with_skew_registry_partner": len(giz & tw_src["skew"]),
                         "skew_hromadas_total": len(tw_src["skew"] & base),
                         "base_rate_skew_pct": round(100 * len(tw_src["skew"] & base) / len(base), 1)}

    # themes: same classifier on both registers (formal agreements)
    import openpyxl
    wb = openpyxl.load_workbook(ROOT / "data/cache/mss/mss_registry.xlsx", read_only=True, data_only=True)
    reg = [r for r in list(wb[wb.sheetnames[0]].iter_rows(values_only=True))[1:] if r and r[0] is not None]
    # multi-label: an agreement counts once under every theme it matches
    dom, dom_n, dom_none, dom_sample = Counter(), 0, 0, []
    for r in reg:
        ids = classify_intl_themes(str(r[1] or ""), str(r[6] or ""))
        dom_n += 1
        if ids:
            dom.update(set(ids))
        else:
            dom_none += 1
            if len(dom_sample) < 8:
                dom_sample.append(str(r[6] or r[1])[:120])
    eu, eu_n, eu_none = Counter(), 0, 0
    for h in load("intl-agreements.json")["hromadas"]:
        for a in h["agreements"]:
            if a.get("partner_country") in {"UA", None}:
                continue
            eu_n += 1
            if a["theme_ids"]:
                eu.update(set(a["theme_ids"]))
            else:
                eu_none += 1

    def pct(c, n):
        return {k: round(100 * v / n, 1) for k, v in c.most_common()}

    out["themes"] = {
        "note": "multi-label share of ALL agreements; EU side = Law 3668-IX register (formal international agreements, not SKEW/AFCCRE twinning)",
        "domestic_registry": {"n": dom_n, "unclassified": dom_none, "pct": pct(dom, dom_n), "unclassified_sample": dom_sample},
        "intl_agreements": {"n": eu_n, "unclassified": eu_none, "pct": pct(eu, eu_n)},
    }

    json.dump(out, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
